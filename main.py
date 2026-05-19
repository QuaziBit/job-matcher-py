import json
import logging
import os
import threading
from contextlib import asynccontextmanager

from fastapi import Depends, FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import HTMLResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from dotenv import load_dotenv

import aiosqlite
from database import get_db, init_db, get_company_meta, upsert_company_meta, upsert_company_vetting
from scraper import scrape_job, assess_job_text_quality
from company_crawler import crawl_company
from mx_validator import validate_email_domain
from analyzer.company_vetter import vet_company, CACHE_TTL_DAYS
from analyzer.snippet_parser import parse_company_snippet
from analyzer import analyze_match, estimate_salary, extract_salary, _job_has_salary, _ollama_model, BLOCKER_KEYWORDS
from analyzer.llm import _verbose
from analyzer.config import anthropic_model, openai_model, gemini_model
from analyzer.known_models import KNOWN_MODELS
from health import run_health_checks
from utils import build_comparison, clean_text

load_dotenv()

logger = logging.getLogger("main")

# ── App setup ──────────────────────────────────────────────────────────────────

@asynccontextmanager
async def lifespan(app: FastAPI):
    await init_db()
    run_health_checks()
    yield


app = FastAPI(title="Job Matcher", lifespan=lifespan)
app.mount("/static", StaticFiles(directory="ui/static"), name="static")
_UI_DIR = os.path.join(os.path.dirname(__file__), "ui")



# ── Pages ──────────────────────────────────────────────────────────────────────

@app.get("/", response_class=HTMLResponse)
async def index():
    return HTMLResponse(open(os.path.join(_UI_DIR, "index.html"), encoding="utf-8").read())
@app.get("/job/{job_id}", response_class=HTMLResponse)
async def job_detail(job_id: int):
    return HTMLResponse(open(os.path.join(_UI_DIR, "job_detail.html"), encoding="utf-8").read())


# ── Vetting page ─────────────────────────────────────────────────────────────

@app.get("/vetting", response_class=HTMLResponse)
async def vetting_page():
    return HTMLResponse(open(os.path.join(_UI_DIR, "vetting.html"), encoding="utf-8").read())


@app.get("/api/vetting")
async def get_vetting_data(db: aiosqlite.Connection = Depends(get_db)):
    """Return all jobs grouped by company and by recruiter for the vetting page."""
    try:
        async with db.execute("""
            SELECT
                j.id, j.title, j.company, j.url, j.scraped_at,
                a.status, a.recruiter_name, a.recruiter_email, a.recruiter_phone
            FROM jobs j
            LEFT JOIN applications a ON a.job_id = j.id
            ORDER BY j.company COLLATE NOCASE, j.scraped_at DESC
        """) as cur:
            rows = [dict(r) for r in await cur.fetchall()]
    except Exception as e:
        logger.error(f"\u2717 get_vetting_data DB error: {e}")
        return JSONResponse({"error": "Database error"}, status_code=500)

    # Group by company
    companies = {}
    for r in rows:
        company = (r.get("company") or "Unknown Company").strip() or "Unknown Company"
        if company not in companies:
            companies[company] = {"company": company, "jobs": []}
        companies[company]["jobs"].append({
            "id":              r["id"],
            "title":           r["title"] or "Untitled",
            "url":             r["url"] or "",
            "scraped_at":      r["scraped_at"] or "",
            "status":          r["status"] or "not_applied",
            "recruiter_name":  r["recruiter_name"] or "",
            "recruiter_email": r["recruiter_email"] or "",
            "recruiter_phone": r["recruiter_phone"] or "",
        })

    # Group by recruiter
    recruiters = {}
    for r in rows:
        name  = (r.get("recruiter_name")  or "").strip()
        email = (r.get("recruiter_email") or "").strip()
        key   = email or name or None
        if not key:
            continue
        if key not in recruiters:
            recruiters[key] = {
                "name":      name,
                "email":     email,
                "phone":     (r.get("recruiter_phone") or "").strip(),
                "companies": set(),
                "jobs":      [],
            }
        recruiters[key]["companies"].add(
            (r.get("company") or "Unknown Company").strip() or "Unknown Company"
        )
        recruiters[key]["jobs"].append({
            "id":         r["id"],
            "title":      r["title"] or "Untitled",
            "company":    (r.get("company") or "Unknown Company").strip(),
            "status":     r["status"] or "not_applied",
            "scraped_at": r["scraped_at"] or "",
        })

    # Convert sets to lists for JSON serialization
    recruiters_list = []
    for v in recruiters.values():
        v["companies"] = sorted(v["companies"])
        recruiters_list.append(v)
    recruiters_list.sort(key=lambda r: (r["name"] or r["email"]).lower())

    # Attach cached LLM vetting data to each company
    company_names = list(companies.keys())
    meta_map = {}
    if company_names:
        try:
            placeholders = ",".join(["?"] * len(company_names))
            async with db.execute(
                f"""SELECT company_name,
                           company_url,
                           glassdoor_url, glassdoor_rating, glassdoor_review_count,
                           linkedin_url, linkedin_employee_count, linkedin_founded,
                           bbb_url, bbb_rating,
                           indeed_url, indeed_rating, indeed_review_count,
                           llm_risk_level, llm_assessment,
                           llm_signals, llm_provider, llm_model, llm_assessed_at
                    FROM company_meta WHERE company_name IN ({placeholders})""",
                company_names,
            ) as cur:
                for row in await cur.fetchall():
                    row = dict(row)
                    import json as _json
                    signals = []
                    try:
                        signals = _json.loads(row.get("llm_signals") or "[]")
                    except Exception:
                        pass
                    meta_map[row["company_name"]] = {
                        "company_url":             row.get("company_url") or "",
                        "glassdoor_url":          row.get("glassdoor_url") or "",
                        "glassdoor_rating":        row.get("glassdoor_rating"),
                        "glassdoor_review_count":  row.get("glassdoor_review_count"),
                        "linkedin_url":            row.get("linkedin_url") or "",
                        "linkedin_employee_count": row.get("linkedin_employee_count") or "",
                        "linkedin_founded":        row.get("linkedin_founded") or "",
                        "bbb_url":                 row.get("bbb_url") or "",
                        "bbb_rating":              row.get("bbb_rating") or "",
                        "indeed_url":              row.get("indeed_url") or "",
                        "indeed_rating":           row.get("indeed_rating"),
                        "indeed_review_count":     row.get("indeed_review_count"),
                        "llm_risk_level":          row.get("llm_risk_level"),
                        "llm_assessment":          row.get("llm_assessment"),
                        "llm_signals":             signals,
                        "llm_provider":            row.get("llm_provider"),
                        "llm_model":               row.get("llm_model"),
                        "llm_assessed_at":         row.get("llm_assessed_at"),
                    }
        except Exception as e:
            logger.warning(f"→ could not fetch company_meta for vetting: {e}")

    companies_list = list(companies.values())
    for c in companies_list:
        c["meta"] = meta_map.get(c["company"], {})

    return JSONResponse({
        "companies":  companies_list,
        "recruiters": recruiters_list,
    })


@app.get("/resumes", response_class=HTMLResponse)
async def resumes_page():
    return HTMLResponse(open(os.path.join(_UI_DIR, "resumes.html"), encoding="utf-8").read())


# ── API Endpoints ──────────────────────────────────────────────────────────────

@app.get("/api/jobs/list")
async def jobs_list(
    request: Request,
    db: aiosqlite.Connection = Depends(get_db),
    page:     int = 1,
    per_page: int = 25,
    search:     str = "",
    status:     str = "",
    score:      str = "",
    provider:   str = "",
    added_days: str = "",
    date_from:  str = "",
    date_to:    str = "",
):
    logger_jl = logging.getLogger("jobs_list")

    if page < 1:
        return JSONResponse({"error": f"invalid page {page!r} — must be a positive integer"}, status_code=400)
    if per_page < 0:
        return JSONResponse({"error": f"invalid per_page {per_page!r} — must be 0 (all) or a positive integer"}, status_code=400)

    valid_statuses = {"", "not_applied", "applied", "interviewing", "offered", "rejected"}
    if status not in valid_statuses:
        return JSONResponse({"error": f"invalid status {status!r} — must be one of: not_applied, applied, interviewing, offered, rejected"}, status_code=400)

    valid_scores = {"", "0", "1", "2", "3", "4", "5"}
    if score not in valid_scores:
        return JSONResponse({"error": f"invalid score {score!r} — must be one of: 0, 1, 2, 3, 4, 5"}, status_code=400)

    valid_providers = {"", "anthropic", "openai", "gemini", "ollama", "manual"}
    if provider not in valid_providers:
        return JSONResponse({"error": f"invalid provider {provider!r} — must be one of: anthropic, openai, gemini, ollama, manual"}, status_code=400)

    logger_jl.info(f"→ /api/jobs/list page={page} per_page={per_page} search={search!r} status={status!r} score={score!r} provider={provider!r} added_days={added_days!r} date_from={date_from!r} date_to={date_to!r}")

    where = []
    args  = []

    if search:
        where.append(
            "(LOWER(j.title) LIKE ? OR LOWER(j.company) LIKE ?"
            " OR LOWER(COALESCE(a.recruiter_name,'')) LIKE ?"
            " OR LOWER(COALESCE(a.recruiter_email,'')) LIKE ?"
            " OR LOWER(COALESCE(a.recruiter_phone,'')) LIKE ?)"
        )
        like = f"%{search.lower()}%"
        args += [like, like, like, like, like]

    if status:
        where.append("COALESCE(a.status, 'not_applied') = ?")
        args.append(status)

    if provider:
        if provider == "manual":
            where.append("j.url LIKE 'manual://%'")
        else:
            where.append(
                "(SELECT llm_provider FROM analyses WHERE job_id = j.id ORDER BY created_at DESC LIMIT 1) = ?"
            )
            args.append(provider)

    if score:
        if score == "0":
            where.append("(SELECT COUNT(*) FROM analyses WHERE job_id = j.id) = 0")
        elif score == "5":
            where.append(
                "COALESCE((SELECT adjusted_score FROM analyses WHERE job_id = j.id ORDER BY created_at DESC LIMIT 1),"
                "(SELECT score FROM analyses WHERE job_id = j.id ORDER BY created_at DESC LIMIT 1), 0) = 5"
            )
        else:
            min_score = int(score)
            where.append(
                "COALESCE((SELECT adjusted_score FROM analyses WHERE job_id = j.id ORDER BY created_at DESC LIMIT 1),"
                "(SELECT score FROM analyses WHERE job_id = j.id ORDER BY created_at DESC LIMIT 1), 0) >= ?"
            )
            args.append(min_score)

    # Simple date filter (added_days)
    if added_days:
        try:
            days = int(added_days)
            where.append("j.scraped_at >= datetime('now', ? || ' days')")
            args.append(f"-{days}")
        except ValueError:
            pass

    # Advanced date range filter
    if date_from:
        where.append("date(j.scraped_at) >= ?")
        args.append(date_from)
    if date_to:
        where.append("date(j.scraped_at) <= ?")
        args.append(date_to)

    where_sql = ("WHERE " + " AND ".join(where)) if where else ""

    base_query = f"""
        SELECT j.id, j.url, j.title, j.company, j.location, j.scraped_at,
               COALESCE(a.status, 'not_applied') as status,
               (SELECT score          FROM analyses WHERE job_id = j.id ORDER BY created_at DESC LIMIT 1) as best_score,
               (SELECT adjusted_score FROM analyses WHERE job_id = j.id ORDER BY created_at DESC LIMIT 1) as adjusted_score,
               (SELECT llm_provider   FROM analyses WHERE job_id = j.id ORDER BY created_at DESC LIMIT 1) as provider,
               (SELECT llm_model     FROM analyses WHERE job_id = j.id ORDER BY created_at DESC LIMIT 1) as last_model,
               CASE WHEN (a.recruiter_name IS NOT NULL AND a.recruiter_name != '')
                      OR (a.recruiter_email IS NOT NULL AND a.recruiter_email != '')
                      OR (a.recruiter_phone IS NOT NULL AND a.recruiter_phone != '')
                    THEN 1 ELSE 0 END as has_recruiter
        FROM jobs j
        LEFT JOIN applications a ON a.job_id = j.id
        {where_sql}
        ORDER BY j.scraped_at DESC
    """

    count_query = f"SELECT COUNT(*) FROM jobs j LEFT JOIN applications a ON a.job_id = j.id {where_sql}"
    try:
        async with db.execute(count_query, args) as cur:
            row   = await cur.fetchone()
            total = row[0] if row else 0
    except Exception as e:
        logger_jl.error(f"✗ count query failed: {e}")
        return JSONResponse({"error": "Failed to load jobs from database. Check the terminal for details."}, status_code=500)

    total_pages = 1
    if per_page > 0 and total > 0:
        total_pages = (total + per_page - 1) // per_page

    if page > total_pages and total_pages > 0:
        page = total_pages

    paginated_query = base_query
    paginated_args  = list(args)
    if per_page > 0:
        offset = (page - 1) * per_page
        paginated_query += " LIMIT ? OFFSET ?"
        paginated_args  += [per_page, offset]

    try:
        async with db.execute(paginated_query, paginated_args) as cur:
            rows = await cur.fetchall()
            jobs = [dict(r) for r in rows]
    except Exception as e:
        logger_jl.error(f"✗ jobs query failed: {e}")
        return JSONResponse({"error": "Failed to load jobs from database. Check the terminal for details."}, status_code=500)

    for job in jobs:
        job["is_manual"] = (job.get("url") or "").startswith("manual://")

    logger_jl.info(f"✓ /api/jobs/list total={total} page={page}/{total_pages} returned={len(jobs)}")

    return JSONResponse({
        "jobs":        jobs,
        "total":       total,
        "page":        page,
        "per_page":    per_page,
        "total_pages": total_pages,
    })


@app.get("/jobs/preview", response_class=HTMLResponse)
async def job_preview_page():
    return HTMLResponse(open(os.path.join(_UI_DIR, "job_preview.html"), encoding="utf-8").read())


@app.post("/api/jobs/scrape")
async def scrape_job_preview(url: str = Form(...), db: aiosqlite.Connection = Depends(get_db)):
    """Scrape a URL and return data + warnings — does NOT save to DB."""
    url = url.strip()

    try:
        async with db.execute("SELECT id FROM jobs WHERE url = ?", (url,)) as cur:
            existing = await cur.fetchone()
    except Exception as e:
        logger.error(f"✗ scrape_job_preview DB error checking duplicate: {e}")
        return JSONResponse({"error": "Database error."}, status_code=500)

    if existing:
        return JSONResponse(
            {"error": "This URL has already been added.", "job_id": existing[0]},
            status_code=409,
        )

    try:
        data = await scrape_job(url)
    except ValueError as e:
        return JSONResponse({"error": str(e)}, status_code=422)
    except Exception as e:
        logger.error(f"✗ scrape_job_preview unexpected error: {e}")
        return JSONResponse({"error": "Unexpected error while scraping."}, status_code=500)

    description = clean_text(data["raw_description"])

    # Run warnings on the scraped text
    text_quality    = assess_job_text_quality(description)
    desc_lower      = description.lower()
    blocker_found   = [kw for kw in BLOCKER_KEYWORDS if kw in desc_lower]
    has_warnings    = bool(blocker_found) or text_quality["level"] != "ok"

    return JSONResponse({
        "url":             url,
        "title":           data["title"],
        "company":         data["company"],
        "location":        data["location"],
        "description":     description,
        "blocker_keywords": blocker_found,
        "text_quality":    text_quality,
        "has_warnings":    has_warnings,
    })


@app.post("/api/jobs/save-preview")
async def save_job_preview(
    url:         str = Form(...),
    title:       str = Form(""),
    company:     str = Form(""),
    location:    str = Form(""),
    company_url: str = Form(""),
    description: str = Form(...),
    db: aiosqlite.Connection = Depends(get_db),
):
    """Save a (possibly edited) scraped job to the DB."""
    url         = url.strip()
    title       = title.strip()
    company     = company.strip()
    location    = location.strip()
    company_url = company_url.strip()
    if company_url and not company_url.startswith(("http://", "https://")):
        company_url = ""
    description = clean_text(description.strip())

    if len(description) < 50:
        return JSONResponse({"error": "Description is too short (minimum 50 characters)."}, status_code=422)

    try:
        async with db.execute("SELECT id FROM jobs WHERE url = ?", (url,)) as cur:
            existing = await cur.fetchone()
    except Exception as e:
        logger.error(f"✗ save_job_preview DB error checking duplicate: {e}")
        return JSONResponse({"error": "Database error."}, status_code=500)

    if existing:
        return JSONResponse(
            {"error": "This URL has already been added.", "job_id": existing[0]},
            status_code=409,
        )

    if len(description) > 8000:
        description = description[:8000] + "\n\n[...truncated for analysis]"

    try:
        async with db.execute(
            "INSERT INTO jobs (url, title, company, location, company_url, raw_description) VALUES (?, ?, ?, ?, ?, ?)",
            (url, title, company, location, company_url, description),
        ) as cur:
            job_id = cur.lastrowid
        await db.commit()
    except Exception as e:
        logger.error(f"✗ save_job_preview DB insert error: {e}")
        return JSONResponse({"error": "Failed to save job."}, status_code=500)

    # Sync company_url to company_meta
    if company and company_url:
        try:
            await db.execute(
                """INSERT INTO company_meta (company_name, company_url)
                   VALUES (?, ?)
                   ON CONFLICT(company_name) DO UPDATE SET company_url = excluded.company_url""",
                (company, company_url),
            )
            await db.commit()
        except Exception as e:
            logger.warning(f"✗ save_job_preview company_meta sync failed: {e}")

    logger.info(f"✓ save_job_preview: job {job_id} saved ({len(description)} chars)")
    return JSONResponse({"job_id": job_id, "title": title, "company": company})


@app.post("/api/jobs/add")
async def add_job(url: str = Form(...), db: aiosqlite.Connection = Depends(get_db)):
    """Scrape a job URL and store it."""
    url = url.strip()

    try:
        async with db.execute("SELECT id FROM jobs WHERE url = ?", (url,)) as cur:
            existing = await cur.fetchone()
    except Exception as e:
        logger.error(f"✗ add_job DB error checking duplicate: {e}")
        return JSONResponse({"error": "Database error. Check the terminal for details."}, status_code=500)

    if existing:
        return JSONResponse({"error": "This URL has already been added.", "job_id": existing[0]}, status_code=409)

    try:
        data = await scrape_job(url)
    except ValueError as e:
        return JSONResponse({"error": str(e)}, status_code=422)
    except Exception as e:
        logger.error(f"✗ add_job unexpected scrape error: {e}")
        return JSONResponse({"error": "Unexpected error while scraping. Check the terminal for details."}, status_code=500)

    try:
        async with db.execute(
            "INSERT INTO jobs (url, title, company, location, raw_description) VALUES (?, ?, ?, ?, ?)",
            (url, data["title"], data["company"], data["location"], clean_text(data["raw_description"])),
        ) as cur:
            job_id = cur.lastrowid
        await db.commit()
    except Exception as e:
        logger.error(f"✗ add_job DB insert error: {e}")
        return JSONResponse({"error": "Failed to save job. Check the terminal for details."}, status_code=500)

    return JSONResponse({"job_id": job_id, "title": data["title"], "company": data["company"]})


@app.post("/api/jobs/add-manual")
async def add_job_manual(
    title: str = Form(""),
    company: str = Form(""),
    location: str = Form(""),
    source_url: str = Form(""),
    company_url: str = Form(""),
    description: str = Form(...),
    db: aiosqlite.Connection = Depends(get_db),
):
    """Store a manually pasted job description (no URL scraping)."""
    import hashlib

    description = clean_text(description.strip())
    if len(description) < 50:
        return JSONResponse(
            {"error": "Description is too short (minimum 50 characters)."},
            status_code=422,
        )

    title         = title.strip()   or "Untitled Job"
    company       = company.strip() or ""
    source_url    = source_url.strip()
    company_url   = company_url.strip()
    if company_url and not company_url.startswith(("http://", "https://")):
        company_url = ""

    # Use provided URL if given, otherwise generate a synthetic manual:// URL
    if source_url:
        job_url = source_url
    else:
        slug    = hashlib.md5(description[:200].encode()).hexdigest()[:12]
        job_url = f"manual://{slug}"

    # Duplicate check against both the resolved URL and a synthetic fallback
    slug          = hashlib.md5(description[:200].encode()).hexdigest()[:12]
    synthetic_url = f"manual://{slug}"

    try:
        async with db.execute(
            "SELECT id FROM jobs WHERE url = ? OR (url = ? AND ? = '')",
            (job_url, synthetic_url, source_url),
        ) as cur:
            existing = await cur.fetchone()
    except Exception as e:
        logger.error(f"✗ add_job_manual DB error checking duplicate: {e}")
        return JSONResponse({"error": "Database error. Check the terminal for details."}, status_code=500)

    if existing:
        return JSONResponse(
            {"error": "This job has already been added.", "job_id": existing[0]},
            status_code=409,
        )

    if len(description) > 8000:
        description = description[:8000] + "\n\n[...truncated for analysis]"

    try:
        async with db.execute(
            "INSERT INTO jobs (url, title, company, location, company_url, raw_description) VALUES (?, ?, ?, ?, ?, ?)",
            (job_url, title, company, location.strip(), company_url, description),
        ) as cur:
            job_id = cur.lastrowid
        await db.commit()
    except Exception as e:
        logger.error(f"✗ add_job_manual DB insert error: {e}")
        return JSONResponse({"error": "Failed to save job. Check the terminal for details."}, status_code=500)

    # Sync company_url to company_meta if provided
    if company and company_url:
        try:
            await db.execute(
                """INSERT INTO company_meta (company_name, company_url)
                   VALUES (?, ?)
                   ON CONFLICT(company_name) DO UPDATE SET company_url = excluded.company_url""",
                (company, company_url),
            )
            await db.commit()
        except Exception as e:
            logger.warning(f"✗ add_job_manual company_meta sync failed: {e}")

    return JSONResponse({"job_id": job_id, "title": title, "company": company})


@app.post("/api/jobs/{job_id}/analyze")
async def analyze_job(
    job_id: int,
    resume_id: int = Form(...),
    provider: str = Form("anthropic"),
    analysis_mode: str = Form(""),
    ollama_model: str = Form(""),
    cloud_model: str = Form(""),
    db: aiosqlite.Connection = Depends(get_db),
):
    try:
        async with db.execute("SELECT raw_description FROM jobs WHERE id = ?", (job_id,)) as cur:
            job = await cur.fetchone()
        if not job:
            raise HTTPException(status_code=404, detail="Job not found")

        async with db.execute("SELECT content FROM resumes WHERE id = ?", (resume_id,)) as cur:
            resume = await cur.fetchone()
        if not resume:
            raise HTTPException(status_code=404, detail="Resume not found")
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"✗ analyze_job DB error fetching job/resume: {e}")
        return JSONResponse({"error": "Database error. Check the terminal for details."}, status_code=500)

    try:
        import time as _time
        # Apply analysis_mode, ollama_model, and cloud_model from form if provided
        _valid_modes = {"fast", "standard", "detailed"}
        if analysis_mode and analysis_mode in _valid_modes:
            os.environ["ANALYSIS_MODE"] = analysis_mode
        if ollama_model and provider == "ollama":
            os.environ["OLLAMA_MODEL"] = ollama_model
        if cloud_model:
            _model_env = {
                "anthropic": "ANTHROPIC_MODEL",
                "openai":    "OPENAI_MODEL",
                "gemini":    "GEMINI_MODEL",
            }
            if env_key := _model_env.get(provider):
                os.environ[env_key] = cloud_model
        _start           = _time.monotonic()
        result           = await analyze_match(resume["content"], job["raw_description"], provider)
        duration_seconds = int(_time.monotonic() - _start)
    except ValueError as e:
        return JSONResponse({"error": str(e)}, status_code=422)
    except Exception as e:
        logger.error(f"✗ analyze_job LLM error for job {job_id}: {e}")
        return JSONResponse({"error": "Analysis failed unexpectedly. Check the terminal for details."}, status_code=500)

    matched_v1 = json.dumps([
        s["skill"] if isinstance(s, dict) else s
        for s in result.get("matched_skills", [])
    ])
    missing_v1 = json.dumps([
        {"skill": s["skill"], "severity": s.get("severity", "minor")}
        if isinstance(s, dict) else {"skill": s, "severity": "minor"}
        for s in result.get("missing_skills", [])
    ])

    try:
        await db.execute(
            """INSERT INTO analyses
               (job_id, resume_id, score, adjusted_score, penalty_breakdown,
                matched_skills, missing_skills, reasoning, llm_provider, llm_model,
                matched_skills_v2, missing_skills_v2, suggestions,
                validation_errors, retry_count, used_fallback, duration_seconds,
                analysis_mode)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            (
                job_id, resume_id,
                result["score"],
                result.get("adjusted_score", result["score"]),
                json.dumps(result.get("penalty_breakdown", {})),
                matched_v1, missing_v1,
                result["reasoning"],
                result["llm_provider"],
                result.get("llm_model", ""),
                json.dumps(result.get("matched_skills", [])),
                json.dumps(result.get("missing_skills", [])),
                json.dumps(result.get("suggestions", [])),
                result.get("validation_errors", ""),
                result.get("retry_count", 0),
                1 if result.get("used_fallback") else 0,
                duration_seconds,
                result.get("analysis_mode", os.getenv("ANALYSIS_MODE", "standard")),
            ),
        )
        await db.commit()
    except Exception as e:
        logger.error(f"✗ analyze_job DB insert error for job {job_id}: {e}")
        return JSONResponse({"error": "Failed to save analysis. Check the terminal for details."}, status_code=500)

    return JSONResponse(result)


@app.post("/api/jobs/{job_id}/estimate-salary")
async def estimate_job_salary(
    job_id: int,
    provider: str = Form("anthropic"),
    model: str = Form(""),
    db: aiosqlite.Connection = Depends(get_db),
):
    """Estimate or extract salary for a job using the configured LLM."""
    try:
        async with db.execute(
            "SELECT title, company, location, raw_description, salary_estimate FROM jobs WHERE id = ?",
            (job_id,)
        ) as cur:
            job = await cur.fetchone()
    except Exception as e:
        logger.error(f"✗ estimate_job_salary DB error fetching job {job_id}: {e}")
        return JSONResponse({"error": "Database error. Check the terminal for details."}, status_code=500)

    if not job:
        raise HTTPException(status_code=404, detail="Job not found")

    if job["salary_estimate"]:
        return JSONResponse(json.loads(job["salary_estimate"]))

    SALARY_INCOMPATIBLE_MODELS = ["gemma3"]
    if provider == "ollama":
        current_model = _ollama_model()
        if any(m in current_model for m in SALARY_INCOMPATIBLE_MODELS):
            return JSONResponse(
                {"error": f"{current_model} is not supported for salary estimation. Please switch to Anthropic or llama3.1:8b."},
                status_code=422,
            )

    has_salary = _job_has_salary(job["raw_description"] or "")

    try:
        if has_salary:
            try:
                result = await extract_salary(
                    title=job["title"] or "",
                    company=job["company"] or "",
                    location=job["location"] or "",
                    job_description=job["raw_description"] or "",
                    provider=provider,
                    model=model,
                )
            except ValueError:
                logger.warning(f"→ extract_salary failed for job {job_id}, falling back to estimate_salary")
                result = await estimate_salary(
                    title=job["title"] or "",
                    company=job["company"] or "",
                    location=job["location"] or "",
                    job_description=job["raw_description"] or "",
                    provider=provider,
                    model=model,
                    _skip_salary_check=True,
                )
                result["source"] = "estimated"
        else:
            result = await estimate_salary(
                title=job["title"] or "",
                company=job["company"] or "",
                location=job["location"] or "",
                job_description=job["raw_description"] or "",
                provider=provider,
                model=model,
            )
            result["source"] = "estimated"
    except ValueError as e:
        return JSONResponse({"error": str(e)}, status_code=422)
    except Exception as e:
        logger.error(f"✗ estimate_job_salary LLM error for job {job_id}: {e}")
        return JSONResponse({"error": "Salary estimation failed unexpectedly. Check the terminal for details."}, status_code=500)

    try:
        await db.execute(
            "UPDATE jobs SET salary_estimate = ? WHERE id = ?",
            (json.dumps(result), job_id)
        )
        await db.commit()
    except Exception as e:
        logger.error(f"✗ estimate_job_salary DB update error for job {job_id}: {e}")

    return JSONResponse(result)


@app.delete("/api/jobs/{job_id}/salary-estimate")
async def clear_salary_estimate(job_id: int, db: aiosqlite.Connection = Depends(get_db)):
    """Clear a cached salary estimate so it can be re-run."""
    try:
        await db.execute("UPDATE jobs SET salary_estimate = '' WHERE id = ?", (job_id,))
        await db.commit()
    except Exception as e:
        logger.error(f"✗ clear_salary_estimate DB error for job {job_id}: {e}")
        return JSONResponse({"error": "Failed to clear salary estimate."}, status_code=500)
    return JSONResponse({"ok": True})


@app.post("/api/jobs/{job_id}/application")
async def upsert_application(
    job_id: int,
    status: str = Form("not_applied"),
    recruiter_name: str = Form(""),
    recruiter_email: str = Form(""),
    recruiter_phone: str = Form(""),
    notes: str = Form(""),
    db: aiosqlite.Connection = Depends(get_db),
):
    try:
        await db.execute(
            """INSERT INTO applications (job_id, status, recruiter_name, recruiter_email, recruiter_phone, notes)
               VALUES (?, ?, ?, ?, ?, ?)
               ON CONFLICT(job_id) DO UPDATE SET
                   status=excluded.status,
                   recruiter_name=excluded.recruiter_name,
                   recruiter_email=excluded.recruiter_email,
                   recruiter_phone=excluded.recruiter_phone,
                   notes=excluded.notes,
                   updated_at=CURRENT_TIMESTAMP""",
            (job_id, status, recruiter_name, recruiter_email, recruiter_phone, notes),
        )
        await db.commit()
    except Exception as e:
        logger.error(f"✗ upsert_application DB error for job {job_id}: {e}")
        return JSONResponse({"error": "Failed to save application. Check the terminal for details."}, status_code=500)
    return JSONResponse({"ok": True})


@app.delete("/api/analyses/{analysis_id}")
async def delete_analysis(analysis_id: int, db: aiosqlite.Connection = Depends(get_db)):
    try:
        async with db.execute("SELECT id FROM analyses WHERE id = ?", (analysis_id,)) as cur:
            if not await cur.fetchone():
                raise HTTPException(status_code=404, detail="Analysis not found")
        await db.execute("DELETE FROM analyses WHERE id = ?", (analysis_id,))
        await db.commit()
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"✗ delete_analysis DB error for analysis {analysis_id}: {e}")
        return JSONResponse({"error": "Failed to delete analysis."}, status_code=500)
    return JSONResponse({"ok": True})


@app.delete("/api/jobs/{job_id}")
async def delete_job(job_id: int, db: aiosqlite.Connection = Depends(get_db)):
    try:
        await db.execute("DELETE FROM jobs WHERE id = ?", (job_id,))
        await db.commit()
    except Exception as e:
        logger.error(f"✗ delete_job DB error for job {job_id}: {e}")
        return JSONResponse({"error": "Failed to delete job."}, status_code=500)
    return JSONResponse({"ok": True})


# ── Job email endpoints ──────────────────────────────────────────────────────

@app.get("/api/jobs/{job_id}/email")
async def get_job_email(job_id: int, db: aiosqlite.Connection = Depends(get_db)):
    """Return saved email HTML for a job, or null if none saved."""
    try:
        async with db.execute(
            "SELECT id, raw_html, created_at FROM job_emails WHERE job_id = ?", (job_id,)
        ) as cur:
            row = await cur.fetchone()
    except Exception as e:
        logger.error(f"✗ get_job_email DB error for job {job_id}: {e}")
        return JSONResponse({"error": "Database error"}, status_code=500)
    if not row:
        return JSONResponse({"email": None})
    return JSONResponse({"email": {"id": row[0], "raw_html": row[1], "created_at": row[2]}})


@app.post("/api/jobs/{job_id}/email")
async def save_job_email(
    job_id: int,
    raw_html: str = Form(...),
    db: aiosqlite.Connection = Depends(get_db),
):
    """Save or replace the email HTML for a job."""
    raw_html = raw_html.strip()
    if not raw_html:
        return JSONResponse({"error": "raw_html is required"}, status_code=422)
    try:
        await db.execute(
            """INSERT INTO job_emails (job_id, raw_html)
               VALUES (?, ?)
               ON CONFLICT(job_id) DO UPDATE SET raw_html=excluded.raw_html,
               created_at=CURRENT_TIMESTAMP""",
            (job_id, raw_html),
        )
        await db.commit()
    except Exception as e:
        logger.error(f"✗ save_job_email DB error for job {job_id}: {e}")
        return JSONResponse({"error": "Failed to save email."}, status_code=500)
    logger.info(f"✓ Email saved for job {job_id}")
    return JSONResponse({"ok": True})


@app.delete("/api/jobs/{job_id}/email")
async def delete_job_email(job_id: int, db: aiosqlite.Connection = Depends(get_db)):
    """Delete the saved email for a job."""
    try:
        await db.execute("DELETE FROM job_emails WHERE job_id = ?", (job_id,))
        await db.commit()
    except Exception as e:
        logger.error(f"✗ delete_job_email DB error for job {job_id}: {e}")
        return JSONResponse({"error": "Failed to delete email."}, status_code=500)
    return JSONResponse({"ok": True})


@app.patch("/api/jobs/{job_id}/url")
async def update_job_url(
    job_id: int,
    url: str = Form(""),
    db: aiosqlite.Connection = Depends(get_db),
):
    """Update or clear the source URL of a saved job."""
    import hashlib

    url = url.strip()

    # Validate URL if provided
    if url and not (url.startswith("http://") or url.startswith("https://")):
        return JSONResponse(
            {"error": "URL must start with http:// or https://"},
            status_code=422,
        )

    # Check job exists
    async with db.execute("SELECT id, raw_description FROM jobs WHERE id = ?", (job_id,)) as cur:
        row = await cur.fetchone()
    if not row:
        return JSONResponse({"error": "Job not found."}, status_code=404)

    # If clearing URL, regenerate synthetic manual:// URL
    if not url:
        desc = row[1] or ""
        slug = hashlib.md5(desc[:200].encode()).hexdigest()[:12]
        url  = f"manual://{slug}"

    try:
        await db.execute("UPDATE jobs SET url = ? WHERE id = ?", (url, job_id))
        await db.commit()
    except Exception as e:
        logger.error(f"✗ update_job_url DB error for job {job_id}: {e}")
        return JSONResponse({"error": "Failed to update URL."}, status_code=500)

    logger.info(f"✓ Job {job_id} URL updated to: {url}")
    return JSONResponse({"ok": True, "url": url})


@app.patch("/api/jobs/{job_id}/title")
async def update_job_title(
    job_id: int,
    title: str = Form(""),
    db: aiosqlite.Connection = Depends(get_db),
):
    """Update the title of a saved job."""
    title = title.strip()

    if not title:
        return JSONResponse({"error": "Title cannot be empty."}, status_code=422)

    async with db.execute("SELECT id FROM jobs WHERE id = ?", (job_id,)) as cur:
        row = await cur.fetchone()
    if not row:
        return JSONResponse({"error": "Job not found."}, status_code=404)

    try:
        await db.execute("UPDATE jobs SET title = ? WHERE id = ?", (title, job_id))
        await db.commit()
    except Exception as e:
        logger.error(f"✗ update_job_title DB error for job {job_id}: {e}")
        return JSONResponse({"error": "Failed to update title."}, status_code=500)

    logger.info(f"✓ Job {job_id} title updated to: {title}")
    return JSONResponse({"ok": True, "title": title})


@app.patch("/api/jobs/{job_id}/company")
async def update_job_company(
    job_id: int,
    company: str = Form(""),
    db: aiosqlite.Connection = Depends(get_db),
):
    """Update the company name of a saved job."""
    company = company.strip()

    async with db.execute("SELECT id, company FROM jobs WHERE id = ?", (job_id,)) as cur:
        row = await cur.fetchone()
    if not row:
        return JSONResponse({"error": "Job not found."}, status_code=404)

    old_company = (row[1] or "").strip()

    try:
        await db.execute("UPDATE jobs SET company = ? WHERE id = ?", (company, job_id))
        await db.commit()
    except Exception as e:
        logger.error(f"✗ update_job_company DB error for job {job_id}: {e}")
        return JSONResponse({"error": "Failed to update company."}, status_code=500)

    # Rename company_meta row so vetting/crawl data follows the new name
    if old_company and company and old_company != company:
        try:
            await db.execute(
                "UPDATE company_meta SET company_name = ? WHERE company_name = ?",
                (company, old_company),
            )
            await db.commit()
            logger.info(f"✓ company_meta renamed: {old_company!r} → {company!r}")
        except Exception as e:
            logger.warning(f"✗ company_meta rename failed: {e}")

    logger.info(f"✓ Job {job_id} company updated to: {company!r}")
    return JSONResponse({"ok": True, "company": company})


@app.patch("/api/jobs/{job_id}/location")
async def update_job_location(
    job_id: int,
    location: str = Form(""),
    db: aiosqlite.Connection = Depends(get_db),
):
    """Update the location of a saved job."""
    location = location.strip()

    async with db.execute("SELECT id FROM jobs WHERE id = ?", (job_id,)) as cur:
        row = await cur.fetchone()
    if not row:
        return JSONResponse({"error": "Job not found."}, status_code=404)

    try:
        await db.execute("UPDATE jobs SET location = ? WHERE id = ?", (location, job_id))
        await db.commit()
    except Exception as e:
        logger.error(f"✗ update_job_location DB error for job {job_id}: {e}")
        return JSONResponse({"error": "Failed to update location."}, status_code=500)

    logger.info(f"✓ Job {job_id} location updated to: {location!r}")
    return JSONResponse({"ok": True, "location": location})


@app.patch("/api/jobs/{job_id}/company-url")
async def update_job_company_url(
    job_id: int,
    company_url: str = Form(""),
    db: aiosqlite.Connection = Depends(get_db),
):
    """Update the company URL for a saved job and sync to company_meta."""
    company_url = company_url.strip()
    if company_url and not company_url.startswith(("http://", "https://")):
        return JSONResponse({"error": "company_url must start with http:// or https://"}, status_code=422)

    async with db.execute("SELECT id, company FROM jobs WHERE id = ?", (job_id,)) as cur:
        row = await cur.fetchone()
    if not row:
        return JSONResponse({"error": "Job not found."}, status_code=404)

    job_company = (row[1] or "").strip()

    try:
        await db.execute("UPDATE jobs SET company_url = ? WHERE id = ?", (company_url, job_id))
        await db.commit()
    except Exception as e:
        logger.error(f"✗ update_job_company_url DB error for job {job_id}: {e}")
        return JSONResponse({"error": "Failed to update company URL."}, status_code=500)

    # Sync to company_meta so vetting page sees it too
    if job_company and company_url:
        try:
            await db.execute(
                """INSERT INTO company_meta (company_name, company_url)
                   VALUES (?, ?)
                   ON CONFLICT(company_name) DO UPDATE SET company_url = excluded.company_url""",
                (job_company, company_url),
            )
            await db.commit()
            logger.info(f"✓ company_meta company_url synced for {job_company!r}")
        except Exception as e:
            logger.warning(f"✗ company_meta company_url sync failed: {e}")

    logger.info(f"✓ Job {job_id} company_url updated to: {company_url!r}")
    return JSONResponse({"ok": True, "company_url": company_url})


@app.post("/api/companies/crawl")
async def crawl_company_endpoint(company_name: str = Form("")):
    """
    Crawl 7 sources for a company name and return merged vetting data.
    Results are cached in company_meta for 7 days.
    """
    company_name = company_name.strip()
    if not company_name:
        return JSONResponse({"error": "company_name is required."}, status_code=422)

    # Return cached result if fresh (within 7 days)
    cached = await get_company_meta(company_name)
    if cached:
        from datetime import datetime, timezone
        try:
            raw = cached["crawled_at"].replace("Z", "+00:00")
            crawled_at = datetime.fromisoformat(raw)
            # SQLite CURRENT_TIMESTAMP returns naive strings — treat as UTC
            if crawled_at.tzinfo is None:
                crawled_at = crawled_at.replace(tzinfo=timezone.utc)
            age_days = (datetime.now(timezone.utc) - crawled_at).days
            if age_days < 7:
                logger.info(f"✓ Returning cached company_meta for: {company_name!r}")
                return JSONResponse({"ok": True, "cached": True, **cached})
        except Exception:
            pass

    logger.info(f"→ Crawling company: {company_name!r}")
    data = await crawl_company(company_name)
    await upsert_company_meta(company_name, data)

    result = await get_company_meta(company_name) or {"company_name": company_name}
    logger.info(f"✓ Crawl complete for: {company_name!r} — {len(data)} fields found")
    return JSONResponse({"ok": True, "cached": False, **result})


@app.get("/api/companies/meta")
async def get_company_meta_endpoint(company_name: str = ""):
    """Return cached company_meta for a given company name, or null if not crawled."""
    company_name = company_name.strip()
    if not company_name:
        return JSONResponse({"error": "company_name is required."}, status_code=422)
    row = await get_company_meta(company_name)
    if row is None:
        return JSONResponse({"ok": True, "cached": False, "company_name": company_name})
    return JSONResponse({"ok": True, "cached": True, **row})


@app.post("/api/companies/meta/update")
async def update_company_meta_endpoint(
    company_name:          str  = Form(""),
    glassdoor_rating:      str  = Form(""),
    glassdoor_review_count:str  = Form(""),
    glassdoor_url:         str  = Form(""),
    indeed_rating:         str  = Form(""),
    indeed_review_count:   str  = Form(""),
    indeed_url:            str  = Form(""),
    bbb_rating:            str  = Form(""),
    bbb_url:               str  = Form(""),
    linkedin_url:          str  = Form(""),
    company_url:           str  = Form(""),
):
    """
    Manually update company_meta fields — ratings, review counts, and URLs.
    Only non-empty fields are written. Zero network requests.
    """
    company_name = company_name.strip()
    if not company_name:
        return JSONResponse({"error": "company_name is required."}, status_code=422)

    data = {}

    # Ratings
    if glassdoor_rating.strip():
        try:
            v = float(glassdoor_rating.strip())
            if 1.0 <= v <= 5.0:
                data["glassdoor_rating"] = v
            else:
                return JSONResponse({"error": "glassdoor_rating must be between 1 and 5."}, status_code=422)
        except ValueError:
            return JSONResponse({"error": "glassdoor_rating must be a number."}, status_code=422)

    if glassdoor_review_count.strip():
        try:
            data["glassdoor_review_count"] = int(glassdoor_review_count.strip())
        except ValueError:
            return JSONResponse({"error": "glassdoor_review_count must be an integer."}, status_code=422)

    if indeed_rating.strip():
        try:
            v = float(indeed_rating.strip())
            if 1.0 <= v <= 5.0:
                data["indeed_rating"] = v
            else:
                return JSONResponse({"error": "indeed_rating must be between 1 and 5."}, status_code=422)
        except ValueError:
            return JSONResponse({"error": "indeed_rating must be a number."}, status_code=422)

    if indeed_review_count.strip():
        try:
            data["indeed_review_count"] = int(indeed_review_count.strip())
        except ValueError:
            return JSONResponse({"error": "indeed_review_count must be an integer."}, status_code=422)

    # BBB grade
    if bbb_rating.strip():
        data["bbb_rating"] = bbb_rating.strip().upper()

    # URLs — basic validation
    for field, val in [
        ("glassdoor_url", glassdoor_url),
        ("indeed_url",    indeed_url),
        ("bbb_url",       bbb_url),
        ("linkedin_url",  linkedin_url),
        ("company_url",   company_url),
    ]:
        val = val.strip()
        if val:
            if not val.startswith("http"):
                return JSONResponse({"error": f"{field} must start with http:// or https://"}, status_code=422)
            data[field] = val

    if not data:
        return JSONResponse({"error": "No fields provided to update."}, status_code=422)

    await upsert_company_meta(company_name, data)
    row = await get_company_meta(company_name) or {}

    logger.info(f"✓ update_company_meta: company={company_name!r} fields={list(data.keys())}")
    return JSONResponse({
        "ok":     True,
        "company": company_name,
        "updated": list(data.keys()),
        "meta":   row,
    })


@app.delete("/api/companies/meta")
async def delete_company_meta_endpoint(
    company_name: str,
    db: aiosqlite.Connection = Depends(get_db),
):
    """Delete all company_meta for a company — ratings, URLs, and LLM vetting."""
    company_name = company_name.strip()
    if not company_name:
        return JSONResponse({"error": "company_name is required."}, status_code=422)
    await db.execute("DELETE FROM company_meta WHERE company_name = ?", (company_name,))
    await db.commit()
    logger.info(f"✓ delete_company_meta: company={company_name!r}")
    return JSONResponse({"ok": True, "company": company_name})


@app.post("/api/companies/parse-snippet")
async def parse_company_snippet_endpoint(
    company_name: str = Form(""),
    text:         str = Form(""),
    provider:     str = Form("anthropic"),
    model:        str = Form(""),
):
    """
    Parse a pasted Google search snippet to extract company ratings.
    The user searches Google for "Company reviews site:glassdoor.com OR site:bbb.org"
    and pastes the raw text. The LLM extracts structured rating data.
    No scraping — zero bot detection risk.
    """
    company_name = company_name.strip()
    text         = text.strip()
    provider     = provider.strip().lower() or "anthropic"
    model        = model.strip()

    if not company_name:
        return JSONResponse({"error": "company_name is required."}, status_code=422)
    if not text:
        return JSONResponse({"error": "text is required."}, status_code=422)
    if len(text) > 5000:
        return JSONResponse({"error": "text too long (max 5000 chars)."}, status_code=422)

    try:
        data = await parse_company_snippet(text, provider, model)
    except ValueError as e:
        return JSONResponse({"error": str(e)}, status_code=422)
    except Exception as e:
        logger.error(f"✗ parse_company_snippet({company_name!r}): {e}")
        return JSONResponse({"error": "Parsing failed — please try again."}, status_code=500)

    if not data:
        return JSONResponse({
            "ok":      True,
            "company": company_name,
            "found":   False,
            "message": "No rating data could be extracted from the pasted text.",
        })

    # Merge into company_meta
    await upsert_company_meta(company_name, data)
    row = await get_company_meta(company_name) or {}

    logger.info(f"✓ parse_company_snippet: company={company_name!r} fields={list(data.keys())}")
    return JSONResponse({
        "ok":      True,
        "company": company_name,
        "found":   True,
        "data":    data,
        "meta":    row,
    })


@app.post("/api/companies/vet")
async def vet_company_endpoint(
    company_name: str = Form(""),
    provider:     str = Form("anthropic"),
    model:        str = Form(""),
    force:        str = Form(""),
):
    """
    Run LLM vetting for a company using its crawled metadata.
    If crawl data doesn't exist, auto-crawls first.
    Results are cached in company_meta for 7 days.
    Returns {risk_level, assessment, signals, provider, model, cached}.
    Only company names and public data are sent to the LLM — no PII.
    """
    company_name = company_name.strip()
    if not company_name:
        return JSONResponse({"error": "company_name is required."}, status_code=422)

    provider = provider.strip().lower() or "anthropic"
    model    = model.strip()
    force_rescan = force.strip().lower() in ("1", "true", "yes")

    # Check 7-day cache first (skip if force rescan)
    row = await get_company_meta(company_name)
    if not force_rescan and row and row.get("llm_assessed_at"):
        from datetime import datetime, timedelta
        try:
            assessed_at = datetime.fromisoformat(row["llm_assessed_at"])
            if datetime.utcnow() - assessed_at < timedelta(days=CACHE_TTL_DAYS):
                import json as _json
                signals = []
                try:
                    signals = _json.loads(row.get("llm_signals") or "[]")
                except Exception:
                    pass
                logger.info(f"✓ Returning cached vetting for: {company_name!r}")
                return JSONResponse({
                    "ok":          True,
                    "cached":      True,
                    "company":     company_name,
                    "risk_level":  row.get("llm_risk_level", "unknown"),
                    "assessment":  row.get("llm_assessment", ""),
                    "signals":     signals,
                    "provider":    row.get("llm_provider", provider),
                    "model":       row.get("llm_model", ""),
                })
        except Exception:
            pass

    # Ensure crawl data exists — auto-crawl if missing
    if not row or not row.get("crawled_at"):
        logger.info(f"→ auto-crawling {company_name!r} before vetting")
        try:
            crawl_data = await crawl_company(company_name)
            await upsert_company_meta(company_name, crawl_data)
            row = await get_company_meta(company_name) or {}
        except Exception as e:
            logger.warning(f"→ crawl failed for {company_name!r}: {e} — proceeding with empty meta")
            row = {}

    # Run LLM vetting
    try:
        result = await vet_company(company_name, row or {}, provider, model)
    except ValueError as e:
        logger.error(f"✗ vet_company failed for {company_name!r}: {e}")
        return JSONResponse({"error": str(e)}, status_code=422)
    except Exception as e:
        logger.error(f"✗ vet_company unexpected error for {company_name!r}: {e}")
        return JSONResponse({"error": "Vetting failed — please try again."}, status_code=500)

    # Cache result
    await upsert_company_vetting(
        company_name,
        result["risk_level"],
        result["assessment"],
        result.get("signals", []),
        result["provider"],
        result["model"],
    )

    return JSONResponse({
        "ok":         True,
        "cached":     False,
        "company":    company_name,
        "risk_level": result["risk_level"],
        "assessment": result["assessment"],
        "signals":    result.get("signals", []),
        "provider":   result["provider"],
        "model":      result["model"],
    })


@app.post("/api/email/validate-domain")
async def validate_email_domain_endpoint(
    email: str = Form(""),
    db: aiosqlite.Connection = Depends(get_db),
):
    """
    Check MX records for an email address domain.
    Results are cached in domain_mx_cache for 24 hours.
    Returns {email, domain, valid, has_mx, mx_records, cached, error}.
    """
    email = email.strip()
    if not email:
        return JSONResponse({"error": "email is required."}, status_code=422)

    result = await validate_email_domain(email, db)
    logger.info(
        f"→ MX check: email={email!r} domain={result['domain']!r} "
        f"has_mx={result['has_mx']} cached={result['cached']}"
    )
    if _verbose():
        logger.info(f"→ MX full result: {result}")
    return JSONResponse(result)


@app.get("/api/email/mx-cache")
async def get_mx_cache(db: aiosqlite.Connection = Depends(get_db)):
    """Return all cached MX results as {domain: {has_mx, mx_records}} map."""
    try:
        async with db.execute(
            "SELECT domain, has_mx, mx_records FROM domain_mx_cache"
        ) as cur:
            rows = [dict(r) for r in await cur.fetchall()]
    except Exception as e:
        logger.error(f"✗ get_mx_cache DB error: {e}")
        return JSONResponse({})
    result = {
        r["domain"]: {
            "has_mx":     bool(r["has_mx"]),
            "mx_records": (r["mx_records"] or "").split(",") if r["mx_records"] else [],
            "checked":    True,
        }
        for r in rows
    }
    return JSONResponse(result)


@app.post("/api/resumes/extract")
async def extract_resume_file(file: UploadFile = File(...)):
    """Extract plain text from an uploaded TXT, PDF, or DOCX file."""
    filename = (file.filename or "").lower()
    raw = await file.read()

    try:
        if filename.endswith(".txt"):
            text = raw.decode("utf-8", errors="replace")

        elif filename.endswith(".pdf"):
            from pdfminer.high_level import extract_text_to_fp
            from pdfminer.layout import LAParams
            import io
            out = io.StringIO()
            extract_text_to_fp(io.BytesIO(raw), out, laparams=LAParams(), output_type="text", codec=None)
            text = out.getvalue()

        elif filename.endswith(".docx"):
            from docx import Document
            import io
            doc = Document(io.BytesIO(raw))
            text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())

        else:
            return JSONResponse(
                {"error": "Unsupported file type. Please upload a TXT, PDF, or DOCX file."},
                status_code=422,
            )
    except Exception as e:
        logger.error(f"✗ extract_resume_file error for {filename}: {e}")
        return JSONResponse({"error": f"Failed to extract text from file: {e}"}, status_code=500)

    text = clean_text(text.strip())
    if len(text) < 50:
        return JSONResponse(
            {"error": "Could not extract enough text from the file (minimum 50 characters)."},
            status_code=422,
        )
    return JSONResponse({"text": text, "char_count": len(text)})


@app.post("/api/resumes/add")
async def add_resume(
    label: str = Form(...),
    content: str = Form(...),
    db: aiosqlite.Connection = Depends(get_db),
):
    try:
        async with db.execute(
            "INSERT INTO resumes (label, content) VALUES (?, ?)", (label.strip(), clean_text(content.strip()))
        ) as cur:
            resume_id = cur.lastrowid
        await db.commit()
    except Exception as e:
        logger.error(f"✗ add_resume DB error: {e}")
        return JSONResponse({"error": "Failed to save resume. Check the terminal for details."}, status_code=500)
    return JSONResponse({"resume_id": resume_id, "label": label.strip()})


@app.get("/api/resumes/{resume_id}")
async def get_resume(resume_id: int, db: aiosqlite.Connection = Depends(get_db)):
    """Return full content of a single resume."""
    try:
        async with db.execute(
            "SELECT id, label, content, created_at, LENGTH(content) as char_count FROM resumes WHERE id = ?",
            (resume_id,),
        ) as cur:
            row = await cur.fetchone()
    except Exception as e:
        logger.error(f"✗ get_resume DB error for id={resume_id}: {e}")
        return JSONResponse({"error": "Database error"}, status_code=500)
    if not row:
        return JSONResponse({"error": "Resume not found."}, status_code=404)
    return JSONResponse(dict(row))


@app.delete("/api/resumes/{resume_id}")
async def delete_resume(resume_id: int, db: aiosqlite.Connection = Depends(get_db)):
    try:
        await db.execute("DELETE FROM resumes WHERE id = ?", (resume_id,))
        await db.commit()
    except Exception as e:
        logger.error(f"✗ delete_resume DB error for resume {resume_id}: {e}")
        return JSONResponse({"error": "Failed to delete resume."}, status_code=500)
    return JSONResponse({"ok": True})


@app.get("/api/jobs/{job_id}/description")
async def get_description(job_id: int, db: aiosqlite.Connection = Depends(get_db)):
    try:
        async with db.execute("SELECT raw_description FROM jobs WHERE id = ?", (job_id,)) as cur:
            row = await cur.fetchone()
    except Exception as e:
        logger.error(f"✗ get_description DB error for job {job_id}: {e}")
        return JSONResponse({"error": "Database error."}, status_code=500)
    if not row:
        raise HTTPException(status_code=404)
    return JSONResponse({"description": row["raw_description"]})


# ── Job detail API ───────────────────────────────────────────────────

@app.get("/api/jobs/{job_id}/detail")
async def get_job_detail(job_id: int, db: aiosqlite.Connection = Depends(get_db)):
    """Return all data needed to render the job detail page as JSON."""
    try:
        async with db.execute("SELECT * FROM jobs WHERE id = ?", (job_id,)) as cur:
            job = await cur.fetchone()
    except Exception as e:
        logger.error(f"✗ get_job_detail DB error fetching job {job_id}: {e}")
        return JSONResponse({"error": "Database error"}, status_code=500)

    if not job:
        raise HTTPException(status_code=404, detail="Job not found")
    job = dict(job)

    try:
        async with db.execute("SELECT * FROM applications WHERE job_id = ?", (job_id,)) as cur:
            app_row = await cur.fetchone()
        application = dict(app_row) if app_row else {}

        async with db.execute("""
            SELECT a.*, r.label as resume_label
            FROM analyses a JOIN resumes r ON r.id = a.resume_id
            WHERE a.job_id = ?
            ORDER BY a.created_at DESC
        """, (job_id,)) as cur:
            analyses = [dict(r) for r in await cur.fetchall()]

        async with db.execute("SELECT id, label FROM resumes ORDER BY created_at DESC") as cur:
            resumes = [dict(r) for r in await cur.fetchall()]
    except Exception as e:
        logger.error(f"✗ get_job_detail DB error fetching related data for job {job_id}: {e}")
        return JSONResponse({"error": "Database error"}, status_code=500)

    # Parse JSON fields in analyses
    for analysis in analyses:
        try:
            v2_matched = analysis.get("matched_skills_v2", "[]") or "[]"
            parsed_v2  = json.loads(v2_matched)
            if parsed_v2:
                analysis["matched_skills"] = parsed_v2
            else:
                v1 = json.loads(analysis.get("matched_skills", "[]") or "[]")
                analysis["matched_skills"] = [
                    {"skill": s, "match_type": "exact",
                     "jd_snippet": "", "resume_snippet": "", "category": "other"}
                    for s in v1 if isinstance(s, str)
                ]
        except Exception:
            analysis["matched_skills"] = []

        try:
            v2_missing = analysis.get("missing_skills_v2", "[]") or "[]"
            parsed_v2  = json.loads(v2_missing)
            if parsed_v2:
                analysis["missing_skills"] = parsed_v2
            else:
                v1 = json.loads(analysis.get("missing_skills", "[]") or "[]")
                analysis["missing_skills"] = [
                    {"skill": s["skill"] if isinstance(s, dict) else s,
                     "severity": s.get("severity", "minor") if isinstance(s, dict) else "minor",
                     "requirement_type": s.get("requirement_type", "preferred") if isinstance(s, dict) else "preferred",
                     "jd_snippet": "", "category": "other"}
                    for s in v1
                ]
        except Exception:
            analysis["missing_skills"] = []

        try:
            pb = analysis.get("penalty_breakdown")
            if isinstance(pb, str):
                analysis["penalty_breakdown"] = json.loads(pb)
        except Exception:
            analysis["penalty_breakdown"] = {}

        try:
            sugg = analysis.get("suggestions", "[]") or "[]"
            analysis["suggestions"] = json.loads(sugg) if isinstance(sugg, str) else sugg
        except Exception:
            analysis["suggestions"] = []

        if not analysis.get("adjusted_score"):
            analysis["adjusted_score"] = analysis["score"]

        analysis["used_fallback"] = bool(analysis.get("used_fallback", 0))

    text_quality   = assess_job_text_quality(job.get("raw_description") or "")
    comparison     = build_comparison(analyses)
    last_resume_id = analyses[0]["resume_id"] if analyses else None

    # company_url: single source of truth is company_meta.
    # If job has a URL but meta doesn't, sync it up first.
    if job.get("company"):
        meta_row = await get_company_meta(job["company"])
        job_url  = (job.get("company_url") or "").strip()
        meta_url = (meta_row.get("company_url") or "").strip() if meta_row else ""
        if job_url and not meta_url:
            try:
                async with db.execute(
                    """INSERT INTO company_meta (company_name, company_url)
                       VALUES (?, ?)
                       ON CONFLICT(company_name) DO UPDATE SET company_url = excluded.company_url""",
                    (job["company"], job_url),
                ) as _:
                    pass
                await db.commit()
                meta_url = job_url
            except Exception as e:
                logger.warning(f"✗ get_job_detail company_meta url sync failed: {e}")
        # Always display from company_meta
        job["company_url"] = meta_url

    salary_data = json.loads(job["salary_estimate"]) if job.get("salary_estimate") else None
    if analyses:
        last_provider = analyses[0]["llm_provider"]
        last_model    = analyses[0].get("llm_model", "")
    elif salary_data and salary_data.get("llm_provider"):
        last_provider = salary_data["llm_provider"]
        last_model    = salary_data.get("llm_model", "")
    else:
        last_provider = "anthropic"
        last_model    = ""

    return JSONResponse({
        "job":              job,
        "application":      application,
        "analyses":         analyses,
        "resumes":          resumes,
        "ollama_model":     last_model if last_provider == "ollama"    else _ollama_model(),
        "anthropic_model":  last_model if last_provider == "anthropic" else anthropic_model(),
        "openai_model":     last_model if last_provider == "openai"    else openai_model(),
        "gemini_model":     last_model if last_provider == "gemini"    else gemini_model(),
        "text_quality":     text_quality,
        "comparison":       comparison,
        "analysis_mode":    analyses[0].get("analysis_mode") or os.getenv("ANALYSIS_MODE", "standard") if analyses else os.getenv("ANALYSIS_MODE", "standard"),
        "salary_estimate":  salary_data,
        "has_salary_in_jd": _job_has_salary(job.get("raw_description", "")),
        "last_resume_id":   last_resume_id,
        "last_provider":    last_provider,
    })


# ── Providers status API ─────────────────────────────────────────────

@app.get("/api/providers/status")
async def get_providers_status():
    """Return which providers are configured and reachable, plus default models."""
    import httpx
    from analyzer.config import ollama_base_url

    has_ollama = False
    try:
        async with httpx.AsyncClient(timeout=2) as client:
            r = await client.get(f"{ollama_base_url()}/api/tags")
            has_ollama = r.status_code == 200
    except Exception:
        pass

    return JSONResponse({
        "has_anthropic":   bool(os.getenv("ANTHROPIC_API_KEY", "")),
        "has_openai":      bool(os.getenv("OPENAI_API_KEY", "")),
        "has_gemini":      bool(os.getenv("GEMINI_API_KEY", "")),
        "has_ollama":      has_ollama,
        "anthropic_model": anthropic_model(),
        "openai_model":    openai_model(),
        "gemini_model":    gemini_model(),
        "ollama_model":    _ollama_model(),
        "mx_auto_check":   os.getenv("MX_AUTO_CHECK", "true").lower() in ("1", "true", "yes"),
    })


# ── Resumes list API ─────────────────────────────────────────────────────────

@app.get("/api/resumes/")
async def list_resumes(db: aiosqlite.Connection = Depends(get_db)):
    """Return all saved resumes for the shared frontend."""
    try:
        async with db.execute(
            "SELECT id, label, created_at, LENGTH(content) as char_count FROM resumes ORDER BY created_at DESC"
        ) as cur:
            resumes = [dict(r) for r in await cur.fetchall()]
    except Exception as e:
        logger.error(f"\u2717 list_resumes DB error: {e}")
        return JSONResponse({"error": "Database error"}, status_code=500)
    return JSONResponse({"resumes": resumes})


# ── Ollama models proxy ──────────────────────────────────────────────────────

@app.get("/api/ollama/models")
async def get_ollama_models():
    """Proxy Ollama /api/tags to avoid CORS issues in the browser."""
    import httpx
    from analyzer.config import ollama_base_url
    try:
        async with httpx.AsyncClient(timeout=3) as client:
            resp = await client.get(f"{ollama_base_url()}/api/tags")
            resp.raise_for_status()
        models = [m["name"] for m in resp.json().get("models", [])]
        return JSONResponse({"models": models})
    except Exception:
        return JSONResponse({"models": []})


# ── Known models endpoint ────────────────────────────────────────────────────

@app.get("/api/providers/models")
async def get_provider_models(provider: str):
    """Return known models for a cloud provider with cost labels."""
    models = KNOWN_MODELS.get(provider, [])
    return JSONResponse({"provider": provider, "models": models})


# ── Entry point ────────────────────────────────────────────────────────────────
if __name__ == "__main__":
    import time
    import uvicorn
    from launcher import Launcher, open_browser

    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s %(message)s",
        datefmt="%Y/%m/%d %H:%M:%S",
    )

    load_dotenv()
    initial_cfg = {
        "port":              int(os.getenv("APP_PORT", "8000")),
        "host":              os.getenv("APP_HOST", "127.0.0.1"),
        "db_path":           os.getenv("DB_PATH", "job_matcher.db"),
        "anthropic_api_key": os.getenv("ANTHROPIC_API_KEY", ""),
        "anthropic_model":   os.getenv("ANTHROPIC_MODEL", ""),
        "openai_api_key":    os.getenv("OPENAI_API_KEY", ""),
        "gemini_api_key":    os.getenv("GEMINI_API_KEY", ""),
        "ollama_base_url":   os.getenv("OLLAMA_BASE_URL", "http://localhost:11434"),
        "ollama_model":      os.getenv("OLLAMA_MODEL", "llama3.1:8b"),
        "ollama_timeout":    int(os.getenv("OLLAMA_TIMEOUT", "600")),
        "analysis_mode":     os.getenv("ANALYSIS_MODE", "standard"),
        "show_more_logs":    os.getenv("SHOW_MORE_LOGS", "").lower() in ("1", "true", "yes"),
        "mx_auto_check":     os.getenv("MX_AUTO_CHECK", "true").lower() in ("1", "true", "yes"),
    }

    launcher_instance = Launcher(initial_cfg)
    uvicorn_server: list = []

    def start_app(cfg: dict):
        host   = cfg.get("host", "127.0.0.1")
        port   = cfg.get("port", 8000)
        key    = cfg.get("anthropic_api_key", "")
        masked = (key[:12] + "...") if key else "not set"

        logger.info(f"\n  Config loaded:")
        logger.info(f"    Anthropic key : {masked}")
        logger.info(f"    Ollama model  : {cfg.get('ollama_model')}")
        logger.info(f"    DB path       : {cfg.get('db_path')}")
        logger.info(f"\n  Starting app on http://{host}:{port} …")

        load_dotenv(override=True)
        os.environ["APP_HOST"]          = host
        os.environ["APP_PORT"]          = str(port)
        os.environ["ANTHROPIC_API_KEY"] = key
        if m := cfg.get("anthropic_model", ""): os.environ["ANTHROPIC_MODEL"] = m
        os.environ["OPENAI_API_KEY"]    = cfg.get("openai_api_key", "")
        os.environ["GEMINI_API_KEY"]    = cfg.get("gemini_api_key", "")
        os.environ["OLLAMA_MODEL"]      = cfg.get("ollama_model", "llama3.1:8b")
        os.environ["OLLAMA_BASE_URL"]   = cfg.get("ollama_base_url", "http://localhost:11434")
        os.environ["OLLAMA_TIMEOUT"]    = str(cfg.get("ollama_timeout", 600))

        config = uvicorn.Config("main:app", host=host, port=port, reload=False)
        server = uvicorn.Server(config)
        uvicorn_server.clear()
        uvicorn_server.append(server)
        threading.Thread(target=server.run, daemon=True).start()
        logger.info(f"✓  Job Matcher running at http://{host}:{port}")

    def stop_app():
        if uvicorn_server:
            logger.info("→ Stopping app server...")
            uvicorn_server[0].should_exit = True
            time.sleep(1)
            uvicorn_server.clear()
            logger.info("✓ App server stopped")
        else:
            logger.info("→ Stop requested but no server running")

    def restart_app(cfg: dict):
        logger.info(f"↺  Restarting app with model={cfg.get('ollama_model')} port={cfg.get('port')}")
        stop_app()
        time.sleep(0.5)
        start_app(cfg)

    launcher_instance.on_start   = start_app
    launcher_instance.on_stop    = stop_app
    launcher_instance.on_restart = restart_app

    launcher_url = launcher_instance.start()

    CYAN  = "\033[96m"
    BOLD  = "\033[1m"
    RESET = "\033[0m"
    print(f"\n{CYAN}{BOLD}{'═' * 54}{RESET}")
    print(f"{CYAN}{BOLD}   Job Matcher{RESET}")
    print(f"{CYAN}{'═' * 54}{RESET}\n")
    print(f"  Launcher   {launcher_url}\n")
    print(f"  Opening launcher in your browser…")
    print(f"  Configure settings and click Start.\n")

    open_browser(launcher_url)

    try:
        while True:
            time.sleep(1)
    except KeyboardInterrupt:
        print("\n  Shutting down…")
        stop_app()
        launcher_instance.stop()
        print("  Stopped. Goodbye!")
